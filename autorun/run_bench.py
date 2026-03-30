#!/usr/bin/env python3
"""
SGLang 自动化压测工具

用法:
    python run_bench.py              # 交互菜单
    python run_bench.py --kill       # 只 kill 当前服务
    python run_bench.py --config /path/to/bench.yaml
"""

import os
import re
import signal
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import yaml

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("[error] 请先安装: pip install openpyxl")

# ── 最重要的指标放最前，用中文列名显示 ──────────────────────────────────
# (原始 key → 列头显示名)
PRIORITY_COLS = [
    ("Mean Input Length",              "平均输入长度 (tok)"),
    ("Mean Output Length",             "平均输出长度 (tok)"),
    ("Request throughput (req/s)",     "QPS (req/s)"),
    ("Output token throughput (tok/s)","TPS (tok/s)"),
    ("Mean Decode",                    "平均解码速度 (tok/s)"),
    ("Mean TTFT (ms)",                 "首token均值时延 (ms)"),
    ("Mean E2EL (ms)",                 "整句均值时延 (ms)"),
]
PRIORITY_KEYS = [k for k, _ in PRIORITY_COLS]
PRIORITY_LABEL = {k: v for k, v in PRIORITY_COLS}

# ── 只保留 Mean，丢弃 Median/P80/P95/P99 的指标组 ────────────────────────
MEAN_ONLY_GROUPS = {"Input Length", "Output Length", "Cached Tokens"}

# ── 完全跳过（绝对时间戳，无比较意义）────────────────────────────────────
SKIP_GROUPS = {"S_TTFT", "S_ITL", "S_E2EL"}

# 其余指标的展示顺序（priority 之后）
EXTRA_ORDER = [
    "Successful requests",
    "Benchmark duration (s)",
    "Total input tokens",
    "Total generated tokens",
    "Total Token throughput (tok/s)",
    "Median Decode", "P80 Decode", "P95 Decode", "P99 Decode", "P99.9 Decode",
    "Median TTFT (ms)", "P80 TTFT (ms)", "P95 TTFT (ms)",
    "P99 TTFT (ms)", "P99.9 TTFT (ms)", "P99.95 TTFT (ms)", "P99.99 TTFT (ms)",
    "Mean TPOT (ms)", "P80 TPOT (ms)", "P95 TPOT (ms)", "P99 TPOT (ms)", "P99.9 TPOT (ms)",
    "Mean ITL (ms)", "P80 ITL (ms)", "P95 ITL (ms)", "P99 ITL (ms)", "P99.9 ITL (ms)",
    "Median E2EL (ms)", "P80 E2EL (ms)", "P95 E2EL (ms)",
    "P99 E2EL (ms)", "P99.9 E2EL (ms)", "P99.95 E2EL (ms)", "P99.99 E2EL (ms)",
]


def should_include(key: str) -> bool:
    """过滤掉不需要的指标行。"""
    # 跳过绝对时间戳组
    for grp in SKIP_GROUPS:
        if grp in key:
            return False
    # Mean-only 组只保留 Mean 开头的
    for grp in MEAN_ONLY_GROUPS:
        if grp in key and not key.startswith("Mean"):
            return False
    return True


# ── 解析 txt 里的 Serving Benchmark Result 段落 ───────────────────────────
def parse_output(text: str) -> dict:
    metrics = {}
    in_summary = False
    for line in text.splitlines():
        stripped = line.strip()
        if "Serving Benchmark Result" in stripped:
            in_summary = True
            continue
        if not in_summary:
            continue
        if re.match(r"^=+$", stripped):          # 段落结束行（50个=）
            break
        m = re.match(r"^(.+?):\s+([\d.]+)\s*$", stripped)
        if m:
            try:
                metrics[m.group(1).strip()] = float(m.group(2))
            except ValueError:
                pass
    return metrics


# ── 等待服务就绪（tail server.log，检测 ready 关键字） ────────────────────
READY_MARKER = "The server is fired up and ready to roll"

def wait_for_server(log_file: Path, timeout: int) -> bool:
    print(f"  等待服务就绪（监听 {log_file.name}），超时 {timeout}s ...", flush=True)
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            text = log_file.read_text(errors="replace")
            if READY_MARKER in text:
                print("  ✓ 服务已就绪", flush=True)
                return True
        except FileNotFoundError:
            pass
        remaining = int(deadline - time.time())
        print(f"  ... 等待启动，剩余 {remaining}s", end="\r", flush=True)
        time.sleep(3)
    print("\n  ✗ 等待超时", flush=True)
    return False


# ── 启动服务（后台） ──────────────────────────────────────────────────────
def start_server(cmd: str, log_file: Path) -> subprocess.Popen:
    with open(log_file, "w") as f:
        proc = subprocess.Popen(
            ["bash", "-c", cmd],
            stdout=f,
            stderr=f,
            preexec_fn=os.setsid,   # 独立进程组，方便批量 kill
        )
    return proc


# ── Kill 服务（端口 + GPU 上所有占用进程） ────────────────────────────────
def kill_server(port: int, proc: subprocess.Popen = None):
    print(f"  kill 服务（端口 {port} + GPU 占用进程）...", flush=True)
    killed = set()

    # 1. 按端口 kill
    r = subprocess.run(f"lsof -ti :{port}", shell=True, capture_output=True, text=True)
    for pid in r.stdout.strip().splitlines():
        pid = pid.strip()
        if pid:
            subprocess.run(f"kill -9 {pid}", shell=True)
            killed.add(pid)

    # 2. 按 GPU 显存占用 kill（nvidia-smi 列出所有使用 GPU 的 PID）
    r = subprocess.run(
        "nvidia-smi --query-compute-apps=pid --format=csv,noheader,nounits",
        shell=True, capture_output=True, text=True,
    )
    for pid in r.stdout.strip().splitlines():
        pid = pid.strip()
        if pid and pid not in killed:
            subprocess.run(f"kill -9 {pid}", shell=True)
            killed.add(pid)

    # 3. 兜底：kill 脚本进程组
    if proc and proc.poll() is None:
        try:
            os.killpg(os.getpgid(proc.pid), signal.SIGKILL)
        except (ProcessLookupError, OSError):
            pass

    if killed:
        print(f"  已 kill PID: {', '.join(sorted(killed))}", flush=True)
    else:
        print("  未发现需要 kill 的进程", flush=True)


# ── 运行压测（前台，tee 到终端+文件） ───────────────────────────────────
def run_infer(cmd: str, log_file: Path, run_dir: Path) -> str:
    print(f"  运行压测，输出实时可见 → {log_file.name}", flush=True)

    # benchmark_serving.py 会把 --save-result 的 JSON 写到 cwd，
    # 这里把 cwd 切到 run_dir，JSON 就直接落在结果目录里
    with open(log_file, "w") as lf:
        proc = subprocess.Popen(
            ["bash", "-c", cmd],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            cwd=str(run_dir),          # ← JSON 落到 run_dir
        )
        lines = []
        for line in proc.stdout:
            sys.stdout.write(line)     # 实时打印到终端
            sys.stdout.flush()
            lf.write(line)             # 同时写 log 文件
            lines.append(line)
        proc.wait()

    return "".join(lines)


# ── 写 Excel ─────────────────────────────────────────────────────────────
def write_excel(all_results: list, out_path: Path):
    # all_results 每项: {name, time, status, server_cmd, infer_cmd, metrics}

    # 收集所有实际出现的 key，过滤后排序
    seen_keys = []
    for r in all_results:
        for k in r.get("metrics", {}):
            if k not in seen_keys and should_include(k):
                seen_keys.append(k)

    # 最终列顺序：priority → extra_order → 剩余
    ordered_keys = []
    for k in PRIORITY_KEYS:
        if k in seen_keys and k not in ordered_keys:
            ordered_keys.append(k)
    for k in EXTRA_ORDER:
        if k in seen_keys and k not in ordered_keys:
            ordered_keys.append(k)
    for k in seen_keys:
        if k not in ordered_keys:
            ordered_keys.append(k)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"

    # 表头：固定列 + 指标列（priority 用中文名，其余用原始 key）
    fixed_cols  = ["实验名称", "运行时间", "状态", "起服务脚本", "起请求脚本"]
    metric_cols = [PRIORITY_LABEL.get(k, k) for k in ordered_keys]
    header = fixed_cols + metric_cols

    header_fill = PatternFill("solid", fgColor="2E75B6")
    header_font = Font(bold=True, color="FFFFFF")
    for ci, h in enumerate(header, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # 数据行
    ok_fill   = PatternFill("solid", fgColor="E2EFDA")
    fail_fill = PatternFill("solid", fgColor="FCE4D6")
    for ri, r in enumerate(all_results, 2):
        status   = r["status"]
        row_fill = ok_fill if status == "ok" else fail_fill

        ws.cell(row=ri, column=1, value=r["name"])
        ws.cell(row=ri, column=2, value=r["time"])
        ws.cell(row=ri, column=3, value=status)
        ws.cell(row=ri, column=4, value=r.get("server_cmd", ""))
        ws.cell(row=ri, column=5, value=r.get("infer_cmd", ""))

        for ci, k in enumerate(ordered_keys, len(fixed_cols) + 1):
            ws.cell(row=ri, column=ci, value=r.get("metrics", {}).get(k))

        for ci in range(1, len(header) + 1):
            ws.cell(row=ri, column=ci).fill = row_fill

    # 自动列宽（脚本列限宽60，其余限宽30）
    for ci, h in enumerate(header, 1):
        col_letter = get_column_letter(ci)
        max_len = max(
            len(str(h)),
            *(len(str(ws.cell(row=ri, column=ci).value or ""))
              for ri in range(2, len(all_results) + 2)),
        )
        limit = 60 if ci in (4, 5) else 30
        ws.column_dimensions[col_letter].width = min(max_len + 2, limit)

    ws.freeze_panes = "F2"   # 冻结到指标列开始
    wb.save(out_path)
    print(f"\n[Excel] 已保存: {out_path}")


# ── 交互菜单 ──────────────────────────────────────────────────────────────
def show_menu(experiments: list) -> list:
    print("\n" + "=" * 55)
    print("  SGLang 自动化压测")
    print("=" * 55)
    print(f"  [0] 全部运行 ({len(experiments)} 个实验)")
    for i, exp in enumerate(experiments, 1):
        print(f"  [{i}] {exp['name']}")
    print("  ─────────────────────────────────────")
    print("  [k] Kill 当前服务")
    print("  [q] 退出")
    print("=" * 55)
    raw = input("选择（多个用逗号，如 1,3）: ").strip().lower()

    if raw in ("q", "quit"):
        sys.exit(0)
    if raw == "k":
        return "kill"
    if raw == "0" or raw == "all":
        return list(range(len(experiments)))
    indices = []
    for part in raw.split(","):
        part = part.strip()
        if part.isdigit():
            idx = int(part) - 1
            if 0 <= idx < len(experiments):
                indices.append(idx)
            else:
                print(f"  [warn] 序号 {part} 超出范围，忽略")
        else:
            print(f"  [warn] 无效输入 '{part}'，忽略")
    return indices


# ── 主流程 ────────────────────────────────────────────────────────────────
def main():
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=Path(__file__).parent / "bench.yaml")
    parser.add_argument("--kill", action="store_true", help="只 kill 当前服务")
    args = parser.parse_args()

    cfg_path = Path(args.config)
    if not cfg_path.exists():
        sys.exit(f"[error] 配置文件不存在: {cfg_path}")

    with open(cfg_path) as f:
        cfg = yaml.safe_load(f)

    g = cfg.get("global", {})
    port            = g.get("port", 3015)
    health_url      = g.get("health_url", f"http://127.0.0.1:{port}/health")
    ready_timeout   = g.get("server_ready_timeout", 300)
    shutdown_wait   = g.get("shutdown_wait", 20)
    results_dir     = Path(cfg_path.parent) / g.get("results_dir", "./results")

    experiments = cfg.get("experiments", [])

    # --kill 快捷操作
    if args.kill:
        kill_server(port)
        return

    # 交互选择
    selection = show_menu(experiments)
    if selection == "kill":
        kill_server(port)
        return
    if not selection:
        print("未选择任何实验，退出。")
        return

    # 创建本次运行目录
    run_id  = datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir = results_dir / run_id
    run_dir.mkdir(parents=True, exist_ok=True)
    print(f"\n结果目录: {run_dir}\n")

    all_results = []

    for idx in selection:
        exp  = experiments[idx]
        name = exp["name"]
        print(f"\n{'─'*55}")
        print(f"[{idx+1}/{len(experiments)}] 实验: {name}")
        print(f"{'─'*55}")

        server_log = run_dir / f"{name}_server.log"
        infer_log  = run_dir / f"{name}_infer.log"

        # 1. 启动服务
        server_cmd = exp.get("server", "").strip()
        proc = None
        if server_cmd:
            print(f"  启动服务 → {server_log.name}", flush=True)
            proc = start_server(server_cmd, server_log)
            ready = wait_for_server(server_log, ready_timeout)
            if not ready:
                print(f"  [SKIP] 服务启动超时，跳过实验 {name}")
                kill_server(port, proc)
                all_results.append({
                    "name": name,
                    "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "status": "server_timeout",
                    "server_cmd": server_cmd,
                    "infer_cmd": infer_cmd if 'infer_cmd' in dir() else "",
                    "metrics": {},
                })
                continue
        else:
            print("  server 为空，跳过启动（假设服务已在运行）")

        # 2. 运行压测
        infer_cmd = exp.get("infer", "").strip()
        output = ""
        if infer_cmd:
            print(f"  压测日志 → {infer_log.name}", flush=True)
            output = run_infer(infer_cmd, infer_log, run_dir)
        else:
            print("  infer 为空，跳过压测")

        # 3. 解析结果
        metrics = parse_output(output)
        status  = "ok" if metrics.get("Successful requests", 0) > 0 else "failed"
        if metrics:
            print(f"  ✓ 解析到 {len(metrics)} 项指标  "
                  f"Req/s={metrics.get('Request throughput (req/s)', 'N/A')}  "
                  f"TTFT_mean={metrics.get('Mean TTFT (ms)', 'N/A')}ms")
        else:
            print("  ✗ 未解析到指标（压测可能失败）")

        all_results.append({
            "name": name,
            "time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "status": status,
            "server_cmd": server_cmd,
            "infer_cmd": infer_cmd,
            "metrics": metrics,
        })

        # 4. Kill 服务
        if server_cmd:
            kill_server(port, proc)
            print(f"  等待 GPU 显存释放 {shutdown_wait}s ...", flush=True)
            time.sleep(shutdown_wait)

    # 5. 写 Excel
    if all_results:
        excel_path = run_dir / f"bench_{run_id}.xlsx"
        write_excel(all_results, excel_path)
    else:
        print("\n没有有效结果，不生成 Excel。")


if __name__ == "__main__":
    main()
